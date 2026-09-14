"""Eşleşme rotasyon dosyası — öğrenci listesi + birim kataloğu içe aktarımı.

Eğitim hemşiresinin sezon başında hazırladığı "AB-CD ROTASYON.xlsx" türü
belgeler: her sayfa bir grup eşleşmesidir (PSÇ = "A-B EŞLEŞMESİ", ÇPC =
"C-D EŞLEŞMESİ") ve üç bölüm taşır:

    satır 1      "A-B EŞLEŞMESİ"  "PAZARTESİ/SALI/ÇARŞAMBA"
    satır 3      # | ADI-SOYADI | EYLÜL … HZN |   | BÖLÜMLER VE SAYILARI | BÖLÜMLER
    satır 4…     öğrenciler (ad hücresi sınıf renginde)  …  yer çizelgesi
    …KURALLARI   birim | "2 ÖĞRENCİ" / "1 KIZ ÖĞRENCİ" / "BU GRUPTA ÖĞRENCİ
                 VERME" | CERRAHİ BRANŞ · YOĞUN BÖLÜM …
    en alt       elle yazılmış serbest kurallar

Buradan YALNIZ öğrenciler ve birimler alınır. Serbest metin kurallar
sistematik ayrıştırılamaz; kullanıcı onları Gereksinimler / Özel Kurallar
ekranlarından kurar.

Genel rotasyon okuyucusu (excel_io.read_rotation_workbook) bu belgeyi
yanlış okuyordu: kural bloğunun satırlarını da öğrenci, "1 ÖĞRENCİ" gibi
metinleri de birim sanıyor, 61 öğrencilik dosyadan 130 öğrenci çıkarıyordu.
Bu modül belgenin yapısını bilerek okur.

Saf mantık: Streamlit'e dokunmaz.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field

import openpyxl

from .excel_io import (
    _find_day_label,
    _find_month_header,
    _theme_colors,
    fill_color_hex,
    is_red_font,
)
from .models import (
    BRANCH_CERRAHI,
    BRANCH_DAHILI,
    BRANCH_DIGER,
    BRANCH_DOGUM_SALONU,
    BRANCH_PALYATIF,
    BRANCH_YOGUN_BAKIM,
    INTENSITY_HAFIF,
    INTENSITY_YOGUN,
    PROGRAM_HEMSIRELIK,
    ClassGroup,
    Student,
    Unit,
)
from .normalize import normalize_key, normalize_unit_name, parse_cell, tr_upper

# Belgede kısaltılmış yazımlar → katalogdaki kanonik ad. Gereksinimler
# birimleri ADLA gösterdiği için bu eşleme kritiktir: "4.KAT DAH" diye
# içe aktarılan birim "4.KAT DAHİLİYE" isteyen gereksinime HİÇ uymaz ve
# bu sessizce olur.
CANONICAL = {
    "2.KAT DOĞUM SERVİSİ": "2.KAT DOĞUM",
    "4.KAT DAH": "4.KAT DAHİLİYE",
    "DOĞ.SAL-Y.DOĞ": "DOĞUM SALONU-YENİ DOĞAN",
    "FTR ÜNİT": "FTR.ÜNİT",
    "KÜÇÜK CERR ODASI": "KÜÇÜK CERRAHİ ODASI",
    "ECZ.SARF": "ECZANE SARF",
}

# Belge yalnız CERRAHİ/DAHİLİ BRANŞ yazıyor; diğerleri ada göre anlaşılır.
BRANCH_BY_NAME = {
    "5.KAT PALYATİF": BRANCH_PALYATIF,
    "7.KAT PALYATİF": BRANCH_PALYATIF,
    "YB/B": BRANCH_YOGUN_BAKIM,
    "YB/C": BRANCH_YOGUN_BAKIM,
    "YB/D": BRANCH_YOGUN_BAKIM,
    "DOĞUM SALONU-YENİ DOĞAN": BRANCH_DOGUM_SALONU,
    "2.KAT DOĞUM": BRANCH_DOGUM_SALONU,
}
EBELIK_PRIORITY = {"DOĞUM SALONU-YENİ DOĞAN", "2.KAT DOĞUM"}
POOL_CANDIDATES = {
    "1.KAT CERR", "3.KAT NOR", "4.KAT DAHİLİYE", "5.KAT PALYATİF",
    "6.KAT CERR", "7.KAT PALYATİF", "ACİL", "FTR.ÜNİT",
    "DİŞ", "LABORATUVAR", "2.KAT ÇOCUK",
}

_PAIRING_RE = re.compile(r"([A-ZÇĞİÖŞÜ])\s*-\s*([A-ZÇĞİÖŞÜ])\s*EŞLEŞME")
_GENDER_TAIL_RE = re.compile(r"^(.*\S)\s+([KE])$")


@dataclass
class PairingStudent:
    full_name: str
    gender: str  # "K" / "E" / "" (belgede ad sonunda " k"/" e" varsa)
    group: int
    color: str  # "#RRGGBB" ya da ""
    class_name: str = ""
    sheet: str = ""
    row: int = 0


@dataclass
class PairingUnit:
    name: str  # kanonik ad
    group: int
    quota: int  # 0 = "BU GRUPTA ÖĞRENCİ VERME"
    gender: str = ""
    branch: str = ""  # yalnız belgede yazıyorsa (cerrahi/dahili)
    intensity: str = ""
    raw_name: str = ""
    quota_text: str = ""
    from_rules: bool = True  # False: yalnız yer çizelgesinde geçiyor


@dataclass
class PairingClass:
    name: str
    color: str
    group: int
    count: int


@dataclass
class PairingResult:
    students: list[PairingStudent] = field(default_factory=list)
    units: list[PairingUnit] = field(default_factory=list)
    classes: list[PairingClass] = field(default_factory=list)
    day_labels: dict[int, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    ignored_assignments: int = 0


# ---------------------------------------------------------------- algılama

def _pairing_letters(ws, header_row: int) -> list[str]:
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                m = _PAIRING_RE.search(tr_upper(v))
                if m:
                    return [m.group(1), m.group(2)]
    return []


def _quota_col(ws, header_row: int):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and "BÖLÜMLER VE SAYILARI" in tr_upper(v):
            return c
    return None


def is_pairing_workbook(path) -> bool:
    """En az bir sayfada "X-Y EŞLEŞMESİ" başlığı + yer çizelgesi var mı."""
    wb = openpyxl.load_workbook(path)
    return any(_sheet_layout(ws) for ws in wb.worksheets)


def _sheet_layout(ws):
    header = _find_month_header(ws)
    if header is None:
        return None
    header_row, month_cols = header
    letters = _pairing_letters(ws, header_row)
    quota_col = _quota_col(ws, header_row)
    if not letters or quota_col is None:
        return None
    return header_row, month_cols, letters, quota_col


# ---------------------------------------------------------------- ayrıştırma

def quota_from_text(text) -> int | None:
    """"2 ÖĞRENCİ" → 2 · "1 KIZ ÖĞRENCİ" → 1 · "…VERME" → 0.

    Sayı yoksa ve metin kişi adlarıysa ("AD SOYAD, AD SOYAD") ad sayısı
    döner. Hiçbiri anlaşılmıyorsa None — çağıran yer çizelgesine düşer.
    """
    if text is None:
        return None
    t = tr_upper(str(text)).strip()
    if not t:
        return None
    if "VERME" in t:
        return 0
    m = re.search(r"\d+", t)
    if m:
        return int(m.group())
    parcalar = [p for p in re.split(r"[,;/]", t) if re.search(r"[A-ZÇĞİÖŞÜ]{2}", p)]
    return len(parcalar) if parcalar else None


def canonical_unit_name(raw: str) -> tuple[str, str]:
    """Hücre metninden (kanonik ad, cinsiyet eki)."""
    p = parse_cell(str(raw))
    temiz = normalize_unit_name(p.base_name)
    return CANONICAL.get(normalize_key(temiz), temiz), p.gender_suffix


def clean_student_name(raw: str) -> tuple[str, str]:
    """Ad hücresi → (temiz ad, cinsiyet).

    Belgede bazı adların sonunda " k" / " e" var — isimden anlaşılmayan
    öğrencilerin cinsiyeti böyle işaretlenmiş. Ayrıca küçük "i" ("KAMiL"),
    çift/son boşluk gibi el yazımı izleri temizlenir.
    """
    s = re.sub(r"\s+", " ", str(raw)).strip()
    s = tr_upper(s)
    m = _GENDER_TAIL_RE.match(s)
    if m:
        return m.group(1).strip(), m.group(2)
    return s, ""


def is_grey(hex_color: str) -> bool:
    """Hücre dolgusu "iptal" grisi mi (#A6A6A6, #BFBFBF, #808080 …).

    Belgede o grupta kullanılmayacak birim, sağdaki BÖLÜMLER listesinde
    griye boyanıyor. Beyaza yakın açık griler (#F2F2F2) sayılmaz: tablo
    zeminlerinde süs olarak da kullanılıyorlar.
    """
    if not hex_color or len(hex_color) != 7:
        return False
    try:
        r, g, b = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    except ValueError:
        return False
    return max(r, g, b) - min(r, g, b) <= 8 and 0x40 <= r <= 0xE0


def _master_col(ws, header_row: int, quota_col: int):
    """Yer çizelgesinin sağındaki "BÖLÜMLER" ana listesi (yoksa None)."""
    for c in range(quota_col + 1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and tr_upper(v).strip() == "BÖLÜMLER":
            return c
    return None


def _is_rules_header(text: str) -> bool:
    """Kural bloğunun başlık satırı mı.

    Belgeler arasında yazım değişiyor: "…A-B EŞLEŞMESİNİN KURALLARI" bir
    dosyada, başka bir dosyada "KURALLARI" kelimesi hiç yok. Öğrenci adında
    geçmesi imkânsız kelimelere bakılır.
    """
    t = tr_upper(text)
    return "KURAL" in t or "EŞLEŞME" in t or " GRUBU" in t


def _read_sheet(ws, theme, group: int, result: PairingResult) -> None:
    header_row, month_cols, letters, quota_col = _sheet_layout(ws)
    name_col = None
    for c in range(1, min(month_cols) + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and "ADI" in tr_upper(v):
            name_col = c
            break
    if name_col is None:
        name_col = min(month_cols) - 1

    label = _find_day_label(ws, header_row)
    if label:
        result.day_labels[group] = label

    # ---- öğrenciler: başlık altından kural bloğuna kadar ----
    ogrenciler: list[PairingStudent] = []
    kural_satiri = None
    for r in range(header_row + 1, ws.max_row + 1):
        ad = ws.cell(r, name_col).value
        if isinstance(ad, str) and _is_rules_header(ad):
            kural_satiri = r
            break
        # Kural bloğu başlıksız başlıyorsa imzası: ay sütununda kontenjan metni
        yan = ws.cell(r, min(month_cols)).value
        if isinstance(yan, str) and quota_from_text(yan) is not None and (
            "ÖĞRENCİ" in tr_upper(yan) or "VERME" in tr_upper(yan)
        ):
            kural_satiri = r
            break
        if not isinstance(ad, str) or not ad.strip():
            continue
        temiz, cinsiyet = clean_student_name(ad)
        if not temiz:
            continue
        if any(ws.cell(r, c).value not in (None, "") for c in month_cols):
            result.ignored_assignments += 1
        ogrenciler.append(
            PairingStudent(
                full_name=temiz, gender=cinsiyet, group=group,
                color=fill_color_hex(ws.cell(r, name_col), theme),
                sheet=ws.title, row=r,
            )
        )

    # ---- sınıflar: renklerin ilk görülme sırası ↔ başlıktaki harf sırası ----
    renkler = [c for c in dict.fromkeys(s.color for s in ogrenciler) if c]
    if len(renkler) == len(letters):
        harf = dict(zip(renkler, letters))
    elif len(renkler) < len(letters) and renkler:
        harf = dict(zip(renkler, letters))
        result.warnings.append(
            f"'{ws.title}': başlık {'-'.join(letters)} eşleşmesi diyor ama "
            f"öğrenci adlarında {len(renkler)} farklı renk var. Sınıfları "
            "kontrol edin."
        )
    else:
        harf = {}
        result.warnings.append(
            f"'{ws.title}': {len(renkler)} farklı ad rengi bulundu, başlıktaki "
            f"{'-'.join(letters)} eşleşmesine oturmuyor — sınıflar boş "
            "bırakıldı, içe aktardıktan sonra elle atayın."
        )
    renksiz = 0
    for s in ogrenciler:
        s.class_name = harf.get(s.color, "")
        renksiz += not s.class_name
    if renksiz and harf:
        result.warnings.append(
            f"'{ws.title}': {renksiz} öğrencinin ad hücresinde sınıf rengi "
            "yok — sınıfı boş kalacak."
        )
    sayim = Counter(s.class_name for s in ogrenciler if s.class_name)
    for renk, ad in harf.items():
        result.classes.append(PairingClass(ad, renk, group, sayim.get(ad, 0)))
    result.students.extend(ogrenciler)

    # ---- yer çizelgesi (BÖLÜMLER VE SAYILARI): kontenjanın ikinci kaynağı ----
    yer: Counter = Counter()
    yer_cinsiyet: dict[str, str] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, quota_col).value
        if not isinstance(v, str) or not v.strip():
            continue
        ad, ek = canonical_unit_name(v)
        if not ad:
            continue
        yer[ad] += 1
        if ek:
            yer_cinsiyet[ad] = ek

    # ---- BÖLÜMLER ana listesinde griye boyananlar: bu grupta iptal ----
    iptal: set[str] = set()
    ana = _master_col(ws, header_row, quota_col)
    if ana is not None:
        for r in range(header_row + 1, ws.max_row + 1):
            h = ws.cell(r, ana)
            if isinstance(h.value, str) and h.value.strip() and is_grey(
                fill_color_hex(h, theme)
            ):
                iptal.add(canonical_unit_name(h.value)[0])

    if yer and sum(yer.values()) != len(ogrenciler):
        result.warnings.append(
            f"'{ws.title}': yer çizelgesinde {sum(yer.values())} satır var, "
            f"öğrenci listesinde {len(ogrenciler)} kişi."
        )

    # ---- kural bloğu: birimlerin yetkili kaynağı ----
    birimler: dict[str, PairingUnit] = {}
    if kural_satiri is not None:
        for r in range(kural_satiri, ws.max_row + 1):
            ad_h = ws.cell(r, name_col).value
            kot_h = ws.cell(r, name_col + 1).value
            if not isinstance(ad_h, str) or kot_h in (None, ""):
                continue
            if _is_rules_header(ad_h):
                continue
            ad, ek = canonical_unit_name(ad_h)
            if not ad:
                continue
            kota = quota_from_text(kot_h)
            yerdeki = yer.get(ad)
            if ad in iptal:
                # Gri boyama açık bir iptal kararı; blokta sayı unutulmuş
                # olabilir (belgede bir kez böyle oldu). İptal kazanır.
                if kota:
                    result.warnings.append(
                        f"'{ws.title}' · {ad}: BÖLÜMLER listesinde griye "
                        f"boyanmış (iptal) ama birim bloğunda hâlâ \"{str(kot_h).strip()}\" "
                        "yazıyor — bu grupta KAPALI aktarıldı."
                    )
                kota = 0
            elif yer and yerdeki is None and kota:
                result.warnings.append(
                    f"'{ws.title}' · {ad}: kural bloğunda {kota} kişi yazıyor "
                    "ama yer çizelgesinde hiç yok — kural bloğundaki sayı "
                    "kullanıldı."
                )
            if kota is None or not re.search(r"\d|VERME", tr_upper(str(kot_h))):
                # sayı yazmıyor (kişi adları vs.) → yer çizelgesi daha güvenilir
                if yerdeki:
                    kota = yerdeki
            elif yerdeki is not None and kota != yerdeki and kota > 0:
                result.warnings.append(
                    f"'{ws.title}' · {ad}: kural bloğu {kota} kişi diyor, yer "
                    f"çizelgesinde {yerdeki} satır var — {kota} kullanıldı."
                )
            if kota is None:
                kota = 1
            ozellik = {
                tr_upper(str(ws.cell(r, c).value)).strip()
                for c in range(name_col + 2, ws.max_column + 1)
                if ws.cell(r, c).value not in (None, "")
            }
            brans = (
                BRANCH_CERRAHI if "CERRAHİ BRANŞ" in ozellik
                else BRANCH_DAHILI if "DAHİLİ BRANŞ" in ozellik
                else ""
            )
            yogunluk = (
                INTENSITY_YOGUN if "YOĞUN BÖLÜM" in ozellik
                else INTENSITY_HAFIF if "HAFİF BÖLÜM" in ozellik
                else ""
            )
            cinsiyet = ek or yer_cinsiyet.get(ad, "")
            if not cinsiyet and is_red_font(ws.cell(r, name_col)):
                result.warnings.append(
                    f"'{ws.title}' · {ad}: kırmızı yazılmış ama k/e eki yok — "
                    "cinsiyet kısıtını kontrol edin."
                )
            birimler[ad] = PairingUnit(
                name=ad, group=group, quota=kota, gender=cinsiyet,
                branch=brans, intensity=yogunluk, raw_name=ad_h.strip(),
                quota_text=str(kot_h).strip(),
            )

    # yalnız yer çizelgesinde geçenler (kural bloğu eksikse)
    for ad, adet in yer.items():
        if ad in iptal and ad not in birimler:
            birimler[ad] = PairingUnit(name=ad, group=group, quota=0,
                                       gender=yer_cinsiyet.get(ad, ""),
                                       from_rules=False)
            result.warnings.append(
                f"'{ws.title}' · {ad}: griye boyanmış (iptal) ama yer "
                "çizelgesinde geçiyor — bu grupta KAPALI aktarıldı."
            )
            continue
        if ad not in birimler:
            birimler[ad] = PairingUnit(
                name=ad, group=group, quota=adet,
                gender=yer_cinsiyet.get(ad, ""), from_rules=False,
            )
            if kural_satiri is not None:
                result.warnings.append(
                    f"'{ws.title}' · {ad}: yer çizelgesinde var ama kural "
                    "bloğunda yok — branş/yoğunluk bilgisi olmadan eklendi."
                )
    result.units.extend(birimler.values())


def read_pairing_workbook(path) -> PairingResult:
    wb = openpyxl.load_workbook(path)
    theme = _theme_colors(wb)
    result = PairingResult()
    grup = 0
    for ws in wb.worksheets:
        if _sheet_layout(ws) is None:
            continue
        grup += 1
        _read_sheet(ws, theme, grup, result)
    if not grup:
        result.warnings.append("Eşleşme sayfası bulunamadı.")
    return result


# ---------------------------------------------------------------- model kurma

def build_units(result: PairingResult, catalog: list[Unit]) -> tuple[list[Unit], list[str]]:
    """Belgedeki birimleri katalog üzerine uygular.

    Belge kontenjan, cinsiyet, aktiflik, branş (cerrahi/dahili) ve yoğunluk
    verir; doldurma politikası ve ebelik önceliği gibi belgede OLMAYAN
    özellikler katalogdan korunur. Katalogda olup belgede geçmeyen birimler
    o grup için PASİF yapılır (belge sezonun gerçeğidir).

    Dönüş: (birimler, bilgi notları).
    """
    notlar: list[str] = []
    katalog = {(normalize_key(u.name), u.group): u for u in catalog}
    gruplar = sorted({u.group for u in result.units})
    cikti: list[Unit] = []
    gorulen: set[tuple[str, int]] = set()

    for pu in result.units:
        anahtar = (normalize_key(pu.name), pu.group)
        gorulen.add(anahtar)
        taban = katalog.get(anahtar)
        if taban is not None:
            u = Unit.from_dict({**taban.to_dict()})
            u.name = taban.name
        else:
            havuz = pu.name in POOL_CANDIDATES
            u = Unit(
                name=pu.name, group=pu.group, quota=1,
                branch_category=BRANCH_BY_NAME.get(pu.name, BRANCH_DIGER),
                pool=havuz, single_slot=not havuz,
                ebelik_priority=pu.name in EBELIK_PRIORITY,
            )
            if catalog:
                notlar.append(f"Grup {pu.group} · {pu.name}: katalogda yok, yeni birim olarak eklendi.")
        u.active = pu.quota > 0
        # kontenjan 0 = "bu grupta öğrenci verme": pasif ama kontenjanı 1
        # tutulur ki kullanıcı aktifleştirince kullanılabilir olsun
        u.quota = pu.quota if pu.quota > 0 else max(1, u.quota)
        u.gender_restriction = pu.gender
        if pu.branch:
            u.branch_category = pu.branch
        if pu.from_rules:
            u.intensity = pu.intensity
        cikti.append(u)

    for (anahtar, g), u in katalog.items():
        if g not in gruplar or (anahtar, g) in gorulen:
            continue
        kapali = Unit.from_dict(u.to_dict())
        kapali.active = False
        cikti.append(kapali)
        if u.active:
            notlar.append(f"Grup {g} · {u.name}: belgede geçmiyor, pasif yapıldı.")

    cikti.sort(key=lambda u: (u.group, u.name))
    return cikti, notlar


def build_students(result: PairingResult, programs: dict[str, str] | None = None):
    """(öğrenciler, sınıflar). programs: sınıf adı → program kodu."""
    programs = programs or {}
    ogrenciler = [
        Student(
            full_name=s.full_name, gender=s.gender, group=s.group,
            class_name=s.class_name,
            program=programs.get(s.class_name, PROGRAM_HEMSIRELIK),
        )
        for s in result.students
    ]
    siniflar = []
    gorulen = set()
    for c in result.classes:
        if c.name in gorulen:
            continue
        gorulen.add(c.name)
        siniflar.append(ClassGroup(name=c.name, color_hex=c.color))
    return ogrenciler, siniflar

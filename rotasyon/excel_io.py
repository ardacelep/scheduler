"""Excel içe/dışa aktarım.

İçe aktarım format-toleranslıdır: sayfa/sütun konumları sabit varsayılmaz,
başlıklar aranarak bulunur (giriş dosyası formatı ileride değişebilir).
Renk kodları okunur: öğrenci adı hücresinin dolgu rengi = sınıf,
kırmızı font + " k"/" e" eki = cinsiyet kısıtı.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .normalize import ParsedCell, normalize_key, parse_cell, tr_upper

# Ay adı varyantları -> kanonik ad (settings.months ile aynı yazım)
MONTH_ALIASES = {
    "EYLÜL": "EYLÜL", "EYLUL": "EYLÜL",
    "EKİM": "EKİM", "EKIM": "EKİM",
    "KASIM": "KASIM",
    "ARALIK": "ARALIK", "ARLK": "ARALIK",
    "OCAK": "OCAK", "OCK": "OCAK",
    "ŞUBAT": "ŞUBAT", "ŞBT": "ŞUBAT", "SUBAT": "ŞUBAT",
    "MART": "MART", "MRT": "MART",
    "NİSAN": "NİSAN", "NSN": "NİSAN", "NISAN": "NİSAN",
    "MAYIS": "MAYIS", "MYS": "MAYIS",
    "HAZİRAN": "HAZİRAN", "HZN": "HAZİRAN", "HAZIRAN": "HAZİRAN",
}

RED_RGB = "FFFF0000"


def _theme_colors(wb) -> list[str]:
    """Tema renk paletini (hex) workbook'un theme XML'inden çıkarır."""
    try:
        xml = wb.loaded_theme
        if isinstance(xml, bytes):
            xml = xml.decode("utf-8")
        # clrScheme içindeki sırayla: dk1, lt1, dk2, lt2, accent1..6, hlink, folHlink
        scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
        if not scheme:
            return []
        colors = []
        for m in re.finditer(
            r"<a:(srgbClr)\s+val=\"([0-9A-Fa-f]{6})\"|<a:(sysClr)[^>]*lastClr=\"([0-9A-Fa-f]{6})\"",
            scheme.group(0),
        ):
            colors.append((m.group(2) or m.group(4)).upper())
        return colors
    except Exception:
        return []


def _apply_tint(hex_color: str, tint: float) -> str:
    def ch(c):
        if tint >= 0:
            return round(c + (255 - c) * tint)
        return round(c * (1 + tint))

    r, g, b = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    return f"{ch(r):02X}{ch(g):02X}{ch(b):02X}"


def fill_color_hex(cell, theme_palette: list[str]) -> str:
    """Hücre dolgu rengini #RRGGBB olarak döndürür; dolgu yoksa ''.

    Excel tema renklerini (theme+tint) ve doğrudan RGB'yi destekler —
    referans dosyada ikisi de kullanılıyor.
    """
    fill = cell.fill
    if fill is None or fill.patternType is None:
        return ""
    fg = fill.fgColor
    if fg is None:
        return ""
    if fg.type == "rgb" and fg.rgb and fg.rgb not in ("00000000",):
        return "#" + fg.rgb[-6:].upper()
    if fg.type == "theme" and theme_palette:
        idx = fg.theme
        # clrScheme sırası ile fill theme indeksi arasında lt1/dk1 takası vardır
        order = [1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11]
        if idx is not None and idx < len(order) and order[idx] < len(theme_palette):
            base = theme_palette[order[idx]]
            tint = fg.tint or 0.0
            return "#" + _apply_tint(base, tint)
    return ""


def is_red_font(cell) -> bool:
    c = cell.font.color if cell.font else None
    return bool(c and c.type == "rgb" and c.rgb == RED_RGB)


@dataclass
class ImportedStudent:
    name: str
    group: int
    class_color: str = ""  # "#RRGGBB" (eski format: renk = sınıf)
    class_text: str = ""  # yeni format: "12/A" gibi metin sınıf sütunu
    student_no: str = ""  # föy no / okul no
    side_note: str = ""  # en soldaki sütun notu ("diş", "diş-geç"...)
    row: int = 0


@dataclass
class ImportedUnit:
    name: str  # cinsiyet eki ayrıştırılmış kanonik aday ad
    group: int
    quota: int = 1
    gender_restriction: str = ""  # "K"/"E"/""


@dataclass
class ImportedAssignment:
    student_name: str
    group: int
    month: str  # kanonik ay adı
    cell: ParsedCell
    red_font: bool = False


@dataclass
class ImportResult:
    students: list[ImportedStudent] = field(default_factory=list)
    units: list[ImportedUnit] = field(default_factory=list)
    assignments: list[ImportedAssignment] = field(default_factory=list)
    class_colors: dict[int, list[str]] = field(default_factory=dict)  # grup -> renkler
    day_labels: dict[int, str] = field(default_factory=dict)  # grup -> "PAZARTESİ/SALI/ÇARŞAMBA"
    warnings: list[str] = field(default_factory=list)


_DAY_WORDS = ("PAZARTESİ", "SALI", "ÇARŞAMBA", "PERŞEMBE", "CUMA")


def _find_month_header(ws: Worksheet) -> tuple[int, dict[int, str]] | None:
    """Ay başlık satırını bulur: (satır no, {sütun: kanonik ay}) döndürür."""
    for r in range(1, min(ws.max_row, 10) + 1):
        months: dict[int, str] = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                key = tr_upper(v.strip())
                if key in MONTH_ALIASES:
                    months[c] = MONTH_ALIASES[key]
        if len(months) >= 5:
            return r, months
    return None


def _find_col(ws: Worksheet, header_row: int, contains: str) -> int | None:
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and contains in tr_upper(v):
            return c
    return None


def _detect_group_no(sheet_name: str, fallback: int) -> int:
    m = re.search(r"([12])\s*\.?\s*GRUP", tr_upper(sheet_name))
    return int(m.group(1)) if m else fallback


def _find_day_label(ws: Worksheet, header_row: int) -> str:
    """Başlık satırının üstünde 'PAZARTESİ/SALI/ÇARŞAMBA' gibi gün etiketi arar."""
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and any(d in tr_upper(v) for d in _DAY_WORDS):
                return v.strip()
    return ""


def _quota_values_to_units(values: list[str], group: int) -> list[ImportedUnit]:
    """Kontenjan listesindeki değerleri birimlere çevirir (n tekrar = n kişi)."""
    counts: dict[tuple[str, str], int] = {}
    names: dict[tuple[str, str], str] = {}
    for v in values:
        key_up = tr_upper(v)
        if "BÖLÜMLER" in key_up or key_up.startswith("TOPLAM"):
            continue  # başlık / toplam satırları
        parsed = parse_cell(v)
        if not parsed.base_name:
            continue
        key = (parsed.key, parsed.gender_suffix)
        counts[key] = counts.get(key, 0) + 1
        names.setdefault(key, parsed.base_name)
    return [
        ImportedUnit(
            name=names[key],
            group=group,
            quota=quota,
            gender_restriction=key[1],
        )
        for key, quota in counts.items()
    ]


def _read_separate_quota_sheet(wb, group_titles: dict[str, int]) -> list[ImportedUnit]:
    """'bölümler-sayılar' gibi ayrı bir sayfadan grup kontenjanlarını okur.

    Sayfada her grubun sütun başlığı, grup sayfasının adıyla (PSÇ/ÇPC)
    eşleşir; altındaki değerler kontenjan listesidir.
    """
    units: list[ImportedUnit] = []
    for ws in wb.worksheets:
        title = tr_upper(ws.title)
        if "BÖLÜM" not in title or "SAYI" not in title:
            continue
        # grup başlıklarının hücrelerini bul
        for r in range(1, min(ws.max_row, 5) + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if not isinstance(v, str):
                    continue
                group = group_titles.get(tr_upper(v.strip()))
                if group is None:
                    continue
                values = [
                    str(ws.cell(rr, c).value)
                    for rr in range(r + 1, ws.max_row + 1)
                    if isinstance(ws.cell(rr, c).value, str)
                    and ws.cell(rr, c).value.strip()
                ]
                units += _quota_values_to_units(values, group)
        break
    return units


def read_rotation_workbook(path: str | Path) -> ImportResult:
    """Referans formattaki (ARDA_ROTASYON.xlsx benzeri) dosyayı okur.

    Ay başlığı içeren her sayfa bir grup rotasyon tablosu sayılır;
    "ÖRN.TEK AY HALİ" gibi tek ay sayfaları (ay sayısı < 5) atlanır.
    """
    wb = openpyxl.load_workbook(path)
    theme = _theme_colors(wb)
    result = ImportResult()
    group_counter = 0
    group_titles: dict[str, int] = {}  # sayfa adı -> grup no (ayrı kontenjan sayfası için)
    groups_without_quota: list[tuple[int, str]] = []

    for ws in wb.worksheets:
        header = _find_month_header(ws)
        if header is None:
            continue
        header_row, month_cols = header
        first_month_col = min(month_cols)
        name_col = _find_col(ws, header_row, "ADI")
        if name_col is None or name_col >= first_month_col:
            # Yeni formatta ad sütununun başlığı boş: ilk ay sütununun solu
            name_col = first_month_col - 1
        if name_col < 1:
            result.warnings.append(
                f"'{ws.title}' sayfasında ad sütunu bulunamadı, sayfa atlandı."
            )
            continue
        group_counter += 1
        group = _detect_group_no(ws.title, group_counter)
        group_titles[tr_upper(ws.title.strip())] = group
        quota_col = _find_col(ws, header_row, "BÖLÜMLER VE SAYILARI")
        class_col = _find_col(ws, header_row, "SINIF")
        no_col = _find_col(ws, header_row, "FÖY") or _find_col(
            ws, header_row, "NO"
        )
        day_label = _find_day_label(ws, header_row)
        if day_label:
            result.day_labels[group] = day_label
        # yapısal sütunlar not/sınıf karışmasın diye ayrı tutulur
        known_cols = {name_col, quota_col, class_col, no_col} | set(month_cols)

        colors_seen: list[str] = []
        sheet_start = len(result.students)  # bu sayfanın öğrenci aralığı
        for r in range(header_row + 1, ws.max_row + 1):
            name = ws.cell(r, name_col).value
            if isinstance(name, str) and name.strip():
                cell = ws.cell(r, name_col)
                color = fill_color_hex(cell, theme)
                if color and color not in colors_seen:
                    colors_seen.append(color)
                class_text = ""
                if class_col is not None:
                    cv = ws.cell(r, class_col).value
                    if cv is not None and str(cv).strip():
                        class_text = str(cv).strip()
                student_no = ""
                if no_col is not None and no_col != class_col:
                    nv = ws.cell(r, no_col).value
                    if nv is not None and str(nv).strip():
                        student_no = str(nv).strip()
                side_note = ""
                for c in range(1, name_col):
                    if c in known_cols:
                        continue
                    v = ws.cell(r, c).value
                    # sıra no gibi sayısal sütunlar not değildir
                    if isinstance(v, str) and v.strip() and not v.strip().isdigit():
                        side_note = v.strip()
                        break
                result.students.append(
                    ImportedStudent(
                        name=name.strip(),
                        group=group,
                        class_color=color,
                        class_text=class_text,
                        student_no=student_no,
                        side_note=side_note,
                        row=r,
                    )
                )
                for c, month in month_cols.items():
                    v = ws.cell(r, c).value
                    if v is None or not str(v).strip():
                        continue
                    result.assignments.append(
                        ImportedAssignment(
                            student_name=name.strip(),
                            group=group,
                            month=month,
                            cell=parse_cell(str(v)),
                            red_font=is_red_font(ws.cell(r, c)),
                        )
                    )
        result.class_colors[group] = colors_seen

        # Sağlamlık: "sınıf" etiketli sütun öğrencilerin yarısından azında
        # doluysa gerçek sınıf sütunu değildir (ör. eski dosyada 'diş' notları
        # yanlışlıkla 'sınıf' başlığı altında). Değerler nota indirgenir.
        sheet_students = result.students[sheet_start:]
        if sheet_students:
            with_class = sum(1 for s in sheet_students if s.class_text)
            if 0 < with_class < len(sheet_students) / 2:
                for s in sheet_students:
                    if s.class_text and not s.side_note:
                        s.side_note = s.class_text
                    s.class_text = ""

        # Kontenjan sütunu: her satır bir kontenjan; aynı ad n kez = n kişilik
        if quota_col is not None:
            values = [
                str(ws.cell(r, quota_col).value)
                for r in range(header_row + 1, ws.max_row + 1)
                if isinstance(ws.cell(r, quota_col).value, str)
                and ws.cell(r, quota_col).value.strip()
            ]
            result.units += _quota_values_to_units(values, group)
        else:
            groups_without_quota.append((group, ws.title))

    # Sayfa içinde kontenjan yoksa ayrı 'bölümler-sayılar' sayfasına bak
    if groups_without_quota:
        extra = _read_separate_quota_sheet(wb, group_titles)
        covered = {u.group for u in extra}
        result.units += extra
        for group, title in groups_without_quota:
            if group not in covered:
                result.warnings.append(
                    f"'{title}' sayfasında ve ayrı kontenjan sayfasında "
                    "'BÖLÜMLER VE SAYILARI' bulunamadı; kontenjanlar "
                    "atamalardan çıkarılamadı."
                )

    if group_counter == 0:
        result.warnings.append(
            "Dosyada ay başlıkları içeren rotasyon sayfası bulunamadı."
        )
    return result


# ======================= ÖĞRENCİ LİSTESİ (SEZON BAŞI) =======================


@dataclass
class RosterStudent:
    """Sezon başı öğrenci listesinden okunan bir satır."""

    name: str
    class_text: str = ""  # "11/A"
    field_text: str = ""  # ALAN/DAL ("Ebe Yardımcılığı"...)
    student_no: str = ""
    school: str = ""
    row: int = 0


def _find_roster_header(ws: Worksheet) -> tuple[int, dict[str, int]] | None:
    """Öğrenci listesi başlık satırını bulur: ad + sınıf sütunu zorunlu."""
    for r in range(1, min(ws.max_row, 10) + 1):
        cols: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            key = tr_upper(v)
            if "ADI" in key and "SOYAD" in key:
                cols["name"] = c
            elif "SINIF" in key:
                cols["class"] = c
            elif "ALAN" in key or "DAL" in key:
                cols["field"] = c
            elif "ÖĞRENCİ NO" in key or key.strip() == "NO":
                cols["no"] = c
            elif "OKUL" in key and "NO" not in key:
                cols["school"] = c
        if "name" in cols and "class" in cols:
            return r, cols
    return None


# Başlıksız "sıra | ad soyad | şube-kod" biçimi (ör. "A-HY", "D-EY").
# Elle hazırlanmış listelerde başlık satırı olmayabilir; şube harfi ve
# program kodu tek hücrede birleşik gelir.
_CODED_ROSTER_RE = re.compile(r"^([A-ZÇĞİÖŞÜ]{1,3})\s*-\s*(HY|SBT|EY|EBE)$")


def _read_coded_roster(ws: Worksheet) -> list[RosterStudent]:
    """Başlıksız 'sıra | ad | şube-kod' (A-HY) biçimini okur.

    Sütunlar: A=sıra no, B=ad soyad, C=şube-program kodu. Kod hücresi
    ``_CODED_ROSTER_RE`` ile eşleşen satırlar öğrenci sayılır; başlık/başlık
    dışı satırlar (kod eşleşmezse) sessizce atlanır.
    """
    out: list[RosterStudent] = []
    for r in range(1, ws.max_row + 1):
        name = ws.cell(r, 2).value
        code = ws.cell(r, 3).value
        if not isinstance(name, str) or not name.strip():
            continue
        if not isinstance(code, str):
            continue
        m = _CODED_ROSTER_RE.match(tr_upper(code.strip()))
        if not m:
            continue
        seq = ws.cell(r, 1).value
        out.append(
            RosterStudent(
                name=re.sub(r"\s+", " ", name).strip(),
                class_text=m.group(1),  # şube harfi
                field_text=m.group(2),  # HY/SBT/EY → program_from_field çözer
                student_no=str(seq).strip() if seq is not None else "",
                row=r,
            )
        )
    return out


def read_student_list_workbook(path: str | Path) -> tuple[list[RosterStudent], list[str]]:
    """Sezon başı öğrenci listesini ('SML ÖĞRENCİ LİSTESİ' benzeri) okur.

    Dönüş: (öğrenciler, uyarılar). Önce başlıklı biçim (ADI SOYADI + SINIF)
    denenir; bulunamazsa başlıksız kodlu biçim (sıra | ad | A-HY). Cinsiyet
    bu listelerde bulunmaz; kullanıcı içe aktarım sonrası Veri Yönetimi'nden
    tamamlar.
    """
    wb = openpyxl.load_workbook(path)
    students: list[RosterStudent] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        header = _find_roster_header(ws)
        if header is None:
            continue
        header_row, cols = header
        for r in range(header_row + 1, ws.max_row + 1):
            name = ws.cell(r, cols["name"]).value
            if not isinstance(name, str) or not name.strip():
                continue
            def _text(key: str) -> str:
                c = cols.get(key)
                v = ws.cell(r, c).value if c else None
                return str(v).strip() if v is not None else ""

            students.append(
                RosterStudent(
                    name=name.strip(),
                    class_text=_text("class"),
                    field_text=_text("field"),
                    student_no=_text("no"),
                    school=_text("school"),
                    row=r,
                )
            )
        break  # ilk uygun sayfa yeterli
    if not students:  # başlık yok — kodlu biçimi dene
        for ws in wb.worksheets:
            coded = _read_coded_roster(ws)
            if coded:
                students = coded
                break
    if not students:
        warnings.append(
            "Dosyada öğrenci listesi bulunamadı (ADI SOYADI + SINIF sütunlu "
            "ya da 'sıra | ad | şube-kod' biçiminde bir sayfa gerekli)."
        )
    return students, warnings


def detect_workbook_kind(path: str | Path) -> str:
    """Dosya türünü sezer: 'rotation' | 'student_list' | 'unknown'.

    Streamlit yüklemeleri için dosya benzeri nesneler de kabul edilir.
    """
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        if _find_month_header(ws) is not None:
            return "rotation"
    for ws in wb.worksheets:
        if _find_roster_header(ws) is not None:
            return "student_list"
    for ws in wb.worksheets:  # başlıksız kodlu liste (A-HY)
        if _read_coded_roster(ws):
            return "student_list"
    return "unknown"


# ======================= DIŞA AKTARIM =======================

from io import BytesIO  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

from .models import Assignment, ClassGroup, Student, Unit  # noqa: E402

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOLD = Font(bold=True)
_RED_FONT = Font(color="FF0000")
_CENTER = Alignment(horizontal="center", vertical="center")


def _cell_value_for(unit: Unit, day_note: str) -> str:
    # Okunur biçim: "FTR.ÜNİT (Erkek)" — içe aktarımda parse_cell geri okur.
    value = unit.name_with_gender
    if day_note:
        value += f" ({day_note})"
    return value


def export_yearly_workbook(
    students: list[Student],
    units: list[Unit],
    classes: list[ClassGroup],
    assignments: list[Assignment],
    months: list[str],
    day_labels: dict[str, str] | None = None,
) -> bytes:
    """Yıllık grup görünümü (madde 14).

    Gerçek belgelerdeki düzen izlenir: sınıf | föy no | sıra | ADI-SOYADI |
    aylar... | BÖLÜMLER VE SAYILARI. Öğrenci adı hücresi sınıf rengiyle dolu,
    cinsiyet kısıtlı birimler kırmızı + k/e ekli.
    """
    day_labels = day_labels or {}
    units_by_id = {u.id: u for u in units}
    class_colors = {c.name: c.color_hex.lstrip("#") for c in classes if c.color_hex}
    cell_map: dict[tuple[str, str], Assignment] = {
        (a.student_id, a.month): a for a in assignments
    }
    NAME_COL = 4
    first_month_col = NAME_COL + 1

    wb = Workbook()
    wb.remove(wb.active)
    groups = sorted({s.group for s in students})
    for group in groups:
        ws = wb.create_sheet(f"{group}. GRUP")
        group_students = sorted(
            [s for s in students if s.group == group],
            key=lambda s: (s.class_name, s.full_name),
        )
        # 1. satır: gün etiketi; 2. satır: sütun başlıkları
        label = day_labels.get(str(group), "")
        if label:
            ws.cell(1, first_month_col, label).font = _BOLD
        header_row = 2
        for c, title in ((1, "sınıf"), (2, "föy no"), (3, "sn"), (4, "ADI-SOYADI")):
            ws.cell(header_row, c, title).font = _BOLD
        for j, m in enumerate(months, start=first_month_col):
            c = ws.cell(header_row, j, m)
            c.font = _BOLD
            c.alignment = _CENTER
        quota_col = first_month_col + len(months) + 1
        ws.cell(header_row, quota_col, "BÖLÜMLER VE SAYILARI").font = _BOLD

        for i, s in enumerate(group_students, start=header_row + 1):
            ws.cell(i, 1, s.class_name)
            ws.cell(i, 2, s.student_no)
            ws.cell(i, 3, i - header_row)  # sıra no
            name_cell = ws.cell(i, NAME_COL, s.full_name)
            color = class_colors.get(s.class_name)
            if color:
                name_cell.fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid"
                )
            name_cell.border = _BORDER
            for j, m in enumerate(months, start=first_month_col):
                a = cell_map.get((s.id, m))
                cell = ws.cell(i, j)
                cell.border = _BORDER
                if a and a.unit_id in units_by_id:
                    u = units_by_id[a.unit_id]
                    cell.value = _cell_value_for(u, a.day_note)
                    if u.gender_restriction:
                        cell.font = _RED_FONT

        # kontenjan listesi: aktif birimler, kontenjan kadar tekrar
        row = header_row + 1
        for u in sorted(
            [u for u in units if u.active and u.group in (0, group) and u.quota],
            key=lambda u: u.name,
        ):
            for _ in range(u.quota):
                c = ws.cell(row, quota_col, u.name_with_gender)
                if u.gender_restriction:
                    c.font = _RED_FONT
                row += 1

        ws.column_dimensions[ws.cell(2, NAME_COL).column_letter].width = 24
        for j in range(first_month_col, first_month_col + len(months)):
            ws.column_dimensions[ws.cell(2, j).column_letter].width = 20
        ws.column_dimensions[ws.cell(2, quota_col).column_letter].width = 26

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_single_month_workbook(
    students: list[Student],
    units: list[Unit],
    assignments: list[Assignment],
    month: str,
    day_labels: dict[str, str] | None = None,
) -> bytes:
    """Tek ay görünümü (madde 15): gruplar yan yana, RENKSİZ sade tablo.

    Başlıkta grubun sınıfları ve haftalık gün etiketi yer alır
    (örn. "B-C Pazartesi-Salı-Çarşamba").
    """
    day_labels = day_labels or {}
    units_by_id = {u.id: u for u in units}
    month_assignments = {
        a.student_id: a for a in assignments if a.month == month
    }

    wb = Workbook()
    ws = wb.active
    ws.title = month
    groups = sorted({s.group for s in students})
    col = 1
    for group in groups:
        group_students = [
            s for s in sorted(students, key=lambda s: s.full_name)
            if s.group == group and s.id in month_assignments
        ]
        # birim adına göre sırala (gerçek belgeler böyle)
        group_students.sort(
            key=lambda s: units_by_id[month_assignments[s.id].unit_id].display_name
            if month_assignments[s.id].unit_id in units_by_id
            else ""
        )
        class_names = "-".join(
            sorted({s.class_name for s in group_students if s.class_name})
        )
        # düzen: sınıf | no | ADI-SOYADI | AY  (gerçek belgelerdeki gibi)
        ws.cell(1, col, day_labels.get(str(group), f"{group}. GRUP")).font = _BOLD
        if class_names:
            ws.cell(2, col, class_names).font = _BOLD
        for j, title in enumerate(("sınıf", "no", "ADI-SOYADI", month)):
            ws.cell(3, col + j, title).font = _BOLD
        for i, s in enumerate(group_students, start=4):
            a = month_assignments[s.id]
            ws.cell(i, col, s.class_name).border = _BORDER
            ws.cell(i, col + 1, s.student_no).border = _BORDER
            ws.cell(i, col + 2, s.full_name).border = _BORDER
            u = units_by_id.get(a.unit_id)
            cell = ws.cell(
                i, col + 3, _cell_value_for(u, a.day_note) if u else ""
            )
            cell.border = _BORDER  # renk yok: madde 15
        ws.column_dimensions[ws.cell(3, col + 2).column_letter].width = 24
        ws.column_dimensions[ws.cell(3, col + 3).column_letter].width = 24
        col += 5  # aradaki boş sütunla birlikte

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

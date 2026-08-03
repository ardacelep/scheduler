"""Dışa aktarımda İŞE YARAMAYAN sütunların olmaması.

Aylık tablodaki "no", yıllık tablodaki "föy no" ve "sn" sütunları çıktıda
yer kaplıyordu ama bilgi taşımıyordu: föy no çoğu sezonda boş, "sn" ise
Excel'in kendi satır numarasının tekrarı.

Önemli ayrım: İÇE AKTARMA tarafı bu sütunları hâlâ OKUYOR — kullanıcının
elle hazırladığı belgelerde varlar. Kaldırılan yalnız bizim ürettiğimiz
çıktı. Alttaki round-trip testi ikisini birden bekçiliyor.
"""
from io import BytesIO

import openpyxl

from rotasyon.excel_io import (
    export_single_month_workbook,
    export_yearly_workbook,
)
from rotasyon.models import Assignment, ClassGroup, Student, Unit

MONTHS = ["EYLÜL", "EKİM"]


def _world(groups=(1,)):
    units, students, asg = [], [], []
    for g in groups:
        u = Unit(name="ACİL", group=g, quota=2)
        units.append(u)
        for i in range(2):
            s = Student(
                full_name=f"AD SOYAD {g}{i}", gender="K", group=g,
                class_name="A", student_no=f"{g}{i}",
            )
            students.append(s)
            asg.append(Assignment(student_id=s.id, month="EYLÜL", unit_id=u.id))
    return students, units, asg


def _yearly(groups=(1,)):
    students, units, asg = _world(groups)
    data = export_yearly_workbook(
        students, units, [ClassGroup(name="A", color_hex="#FFFF00")], asg, MONTHS
    )
    return openpyxl.load_workbook(BytesIO(data))


def _monthly(groups=(1,)):
    students, units, asg = _world(groups)
    data = export_single_month_workbook(students, units, asg, "EYLÜL")
    return openpyxl.load_workbook(BytesIO(data)).active


def _headers(ws, rows=range(1, 4)):
    return {
        ws.cell(r, c).value
        for r in rows for c in range(1, ws.max_column + 1)
    }


# ---- yıllık ----

def test_yearly_drops_foy_no_and_sira():
    ws = _yearly()["1. GRUP"]
    basliklar = _headers(ws)
    assert "föy no" not in basliklar
    assert "sn" not in basliklar


def test_yearly_name_sits_next_to_class():
    ws = _yearly()["1. GRUP"]
    assert ws.cell(2, 1).value == "sınıf"
    assert ws.cell(2, 2).value == "ADI-SOYADI"


def test_yearly_months_start_right_after_the_name():
    ws = _yearly()["1. GRUP"]
    assert [ws.cell(2, c).value for c in (3, 4)] == MONTHS


def test_yearly_student_no_is_not_written_anywhere():
    """Föy no sütunu kalktı; değer başka bir hücreye sızmamalı."""
    ws = _yearly()["1. GRUP"]
    tum = {
        str(ws.cell(r, c).value)
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
    }
    assert "10" not in tum and "11" not in tum


def test_yearly_quota_column_still_there():
    """Karşıt kanıt: işe YARAYAN sütun kaldırılmadı."""
    ws = _yearly()["1. GRUP"]
    assert "BÖLÜMLER VE SAYILARI" in _headers(ws)


# ---- aylık ----

def test_monthly_drops_the_no_column():
    ws = _monthly()
    assert "no" not in _headers(ws)


def test_monthly_layout_is_class_name_month():
    ws = _monthly()
    assert [ws.cell(3, c).value for c in (1, 2, 3)] == [
        "sınıf", "ADI-SOYADI", "EYLÜL"
    ]


def test_monthly_group_blocks_stay_separated():
    """Blok genişliği 5→4 oldu; iki grup arasında boş sütun KALMALI.

    Kalmazsa iki grup tek tabloymuş gibi görünür.
    """
    ws = _monthly(groups=(1, 2))
    ad = [
        c for c in range(1, ws.max_column + 1)
        if ws.cell(3, c).value == "ADI-SOYADI"
    ]
    assert len(ad) == 2
    assert ad[1] - ad[0] == 4  # sınıf|ad|ay + boşluk
    ayirici = ad[0] + 2  # ay sütunundan sonraki sütun
    assert all(
        ws.cell(r, ayirici).value in (None, "")
        for r in range(1, ws.max_row + 1)
    )


# ---- içe aktarma bozulmamalı ----

def test_exported_yearly_can_still_be_reimported():
    """Sütunlar kalkınca içe aktarıcı ad sütununu kaybetmemeli."""
    from rotasyon.excel_io import read_rotation_workbook
    from rotasyon.models import DEFAULT_SETTINGS

    # içe aktarıcı ay başlık satırını en az 5 ay görünce tanıyor
    aylar = DEFAULT_SETTINGS["months"]
    students, units, asg = _world()
    data = export_yearly_workbook(
        students, units, [ClassGroup(name="A", color_hex="#FFFF00")], asg, aylar
    )
    result = read_rotation_workbook(BytesIO(data))
    okunan = {s.name for s in result.students}
    assert okunan == {s.full_name for s in students}, result.warnings


def test_importer_still_reads_foy_no_from_hand_made_files():
    """Kullanıcının kendi belgesinde föy no varsa OKUNMAYA devam etmeli."""
    from openpyxl import Workbook

    from rotasyon.excel_io import read_rotation_workbook

    from rotasyon.models import DEFAULT_SETTINGS

    wb = Workbook()
    ws = wb.active
    ws.title = "1. GRUP"
    for c, t in ((1, "sınıf"), (2, "föy no"), (3, "ADI-SOYADI")):
        ws.cell(2, c, t)
    for j, m in enumerate(DEFAULT_SETTINGS["months"], start=4):
        ws.cell(2, j, m)
    ws.cell(3, 1, "A")
    ws.cell(3, 2, "77")
    ws.cell(3, 3, "AD SOYAD")
    ws.cell(3, 4, "ACİL")
    buf = BytesIO()
    wb.save(buf)
    result = read_rotation_workbook(BytesIO(buf.getvalue()))
    assert [s.student_no for s in result.students] == ["77"]

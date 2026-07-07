"""Excel dışa aktarım testleri (madde 14: renkli yıllık, madde 15: renksiz tek ay)."""
from io import BytesIO

import openpyxl
import pytest

from rotasyon.excel_io import (
    export_single_month_workbook,
    export_yearly_workbook,
    read_rotation_workbook,
)
from rotasyon.importer import build_dataset
from rotasyon.models import DEFAULT_SETTINGS

from conftest import REFERENCE_XLSX

pytestmark = pytest.mark.skipif(
    not REFERENCE_XLSX.exists(), reason="referans xlsx yok"
)


@pytest.fixture(scope="module")
def seeded():
    return build_dataset(read_rotation_workbook(REFERENCE_XLSX))


@pytest.fixture(scope="module")
def yearly_wb(seeded):
    data = export_yearly_workbook(
        seeded.students,
        seeded.units,
        seeded.classes,
        seeded.assignments,
        DEFAULT_SETTINGS["months"],
    )
    return openpyxl.load_workbook(BytesIO(data))


def _find_header(ws, title):
    """Başlık hücresini (satır, sütun) olarak bulur."""
    for r in range(1, 4):
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value == title:
                return r, c
    raise AssertionError(f"başlık bulunamadı: {title}")


def test_rule14_yearly_export_has_class_colors(seeded, yearly_wb):
    assert set(yearly_wb.sheetnames) == {"1. GRUP", "2. GRUP"}
    ws = yearly_wb["1. GRUP"]
    _, name_col = _find_header(ws, "ADI-SOYADI")
    class_colors = {
        c.color_hex.lstrip("#").upper() for c in seeded.classes if c.color_hex
    }
    found = set()
    for row in ws.iter_rows(min_row=3, min_col=name_col, max_col=name_col):
        cell = row[0]
        if cell.value and cell.fill.patternType == "solid":
            found.add(cell.fill.fgColor.rgb[-6:].upper())
    assert found and found <= class_colors


def test_yearly_export_has_class_and_no_columns(seeded, yearly_wb):
    """Gerçek belge düzeni: sınıf ve föy no sütunları."""
    ws = yearly_wb["1. GRUP"]
    _find_header(ws, "sınıf")
    _find_header(ws, "föy no")
    header_row, name_col = _find_header(ws, "ADI-SOYADI")
    # ilk öğrenci satırında sınıf dolu
    assert ws.cell(header_row + 1, 1).value  # sınıf metni


def test_rule14_yearly_export_red_gendered_units(yearly_wb):
    """Cinsiyet kısıtlı birimler okunur '(Kız)'/'(Erkek)' ekiyle ve kırmızı."""
    ws = yearly_wb["1. GRUP"]
    header_row, name_col = _find_header(ws, "ADI-SOYADI")
    red_gendered = []
    for row in ws.iter_rows(
        min_row=header_row + 1, min_col=name_col + 1, max_col=name_col + 10
    ):
        for cell in row:
            if isinstance(cell.value, str) and (
                "(Kız)" in cell.value or "(Erkek)" in cell.value
            ):
                red_gendered.append(
                    cell.font.color is not None
                    and cell.font.color.rgb in ("00FF0000", "FFFF0000")
                )
    assert red_gendered and all(red_gendered)


def test_yearly_export_quota_column(yearly_wb):
    ws = yearly_wb["1. GRUP"]
    header_row, header_col = _find_header(ws, "BÖLÜMLER VE SAYILARI")
    values = [
        ws.cell(r, header_col).value
        for r in range(header_row + 1, ws.max_row + 1)
        if ws.cell(r, header_col).value
    ]
    assert values.count("DİŞ") == 4  # kontenjan kadar tekrar


def test_rule15_single_month_export_is_colorless(seeded):
    data = export_single_month_workbook(
        seeded.students,
        seeded.units,
        seeded.assignments,
        "EYLÜL",
        DEFAULT_SETTINGS["day_labels"],
    )
    ws = openpyxl.load_workbook(BytesIO(data)).active
    assert ws.title == "EYLÜL"
    # başlıklar: iki grup bloğu yan yana, gün etiketleri var
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    texts = [h for h in headers if h]
    assert len(texts) == 2
    assert "Pazartesi" in texts[0]
    # hiçbir hücrede dolgu rengi yok (madde 15)
    for row in ws.iter_rows():
        for cell in row:
            assert cell.fill.patternType is None
    # her iki grup bloğunda ad sütunları dolu (blok düzeni: sınıf|no|ad|ay)
    name_cols = [
        c for c in range(1, ws.max_column + 1)
        if ws.cell(3, c).value == "ADI-SOYADI"
    ]
    assert len(name_cols) == 2
    names = sum(
        1
        for c in name_cols
        for r in range(4, ws.max_row + 1)
        if ws.cell(r, c).value
    )
    assert names == 70


def test_roundtrip_reimport(seeded):
    """Dışa aktarılan yıllık dosya, içe aktarıcı tarafından geri okunabilir."""
    data = export_yearly_workbook(
        seeded.students,
        seeded.units,
        seeded.classes,
        seeded.assignments,
        DEFAULT_SETTINGS["months"],
    )
    result = read_rotation_workbook(BytesIO(data))
    ds2 = build_dataset(result)
    assert len(ds2.students) == len(seeded.students)
    assert len(ds2.assignments) == len(seeded.assignments)
    # cinsiyet kısıtları k/e ekiyle korunur
    restrictions = {
        (u.group, u.name, u.gender_restriction)
        for u in ds2.units
        if u.gender_restriction and u.active
    }
    assert (1, "ALÇI ODASI", "E") in restrictions

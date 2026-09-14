"""Eşleşme rotasyon dosyası içe aktarımı (rotasyon/pairing_import.py).

Testler SENTETİK bir çalışma kitabı üzerinde koşar: gerçek belgeler öğrenci
adları içerdiği için depoda yok. Sentetik dosya gerçek belgelerde görülen
el yazımı izlerini bilerek taşır: küçük "i", çift boşluk, ad sonunda " k",
kontenjan yerine kişi adları, kısaltılmış birim adları, özelliklerin farklı
sütunlarda olması, başlığında "KURALLARI" yazmayan kural bloğu.

Gerçek belgeler (example/, git dışı) varsa yalnız SAYILAR doğrulanır.
"""
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from rotasyon.models import (
    BRANCH_CERRAHI,
    BRANCH_DAHILI,
    INTENSITY_HAFIF,
    INTENSITY_YOGUN,
    PROGRAM_EBELIK,
    PROGRAM_HEMSIRELIK,
    Unit,
)
from rotasyon.pairing_import import (
    build_students,
    build_units,
    canonical_unit_name,
    clean_student_name,
    is_pairing_workbook,
    quota_from_text,
    read_pairing_workbook,
)

AYLAR = ["EYLÜL", "EKİM", "KASIM", "ARLK", "OCK", "ŞBT", "MRT", "NSN", "MYS", "HZN"]
SARI, YESIL, PEMBE, MAVI = "FFF2CC", "92D050", "E5A3E5", "57FFA3"
KIRMIZI = Font(color="FFFF0000")


def _dolgu(renk):
    return PatternFill(start_color=renk, end_color=renk, fill_type="solid")


def _sayfa(wb, baslik, harfler, gunler, ogrenciler, yerler, kurallar,
           kural_basligi="KURALLARI"):
    """ogrenciler: [(ad, renk)] · yerler: [birim hücresi] ·
    kurallar: [(ad, kontenjan_metni, [özellikler], kırmızı_mı)]"""
    ws = wb.create_sheet(baslik)
    ws.cell(1, 2, f"{harfler} EŞLEŞMESİ")
    ws.cell(1, 3, gunler)
    ws.cell(3, 2, "ADI- SOYADI")
    for j, ay in enumerate(AYLAR, start=3):
        ws.cell(3, j, ay)
    ws.cell(3, 14, "BÖLÜMLER VE SAYILARI")
    ws.cell(3, 16, "BÖLÜMLER")
    for i, (ad, renk) in enumerate(ogrenciler, start=4):
        ws.cell(i, 1, i - 3)
        c = ws.cell(i, 2, ad)
        if renk:
            c.fill = _dolgu(renk)
    for i, yer in enumerate(yerler, start=4):
        ws.cell(i, 14, yer)
    r = 4 + max(len(ogrenciler), len(yerler)) + 1
    ws.cell(r, 2, f"PZT GRUBU {harfler} EŞLEŞMESİNİN {kural_basligi}".strip())
    for k, (ad, kot, ozellikler, kirmizi) in enumerate(kurallar, start=r + 1):
        c = ws.cell(k, 2, ad)
        if kirmizi:
            c.font = KIRMIZI
        ws.cell(k, 3, kot)
        for j, oz in enumerate(ozellikler, start=4):
            ws.cell(k, j, oz)
    son = r + 1 + len(kurallar) + 1
    ws.cell(son, 2, "10 AYLIK SEZON SONUNDA HER BİR ÖĞRENCİ BİRER DEFA YB GÖRMÜŞ OLMALI")
    return ws


def _kitap(kural_basligi="KURALLARI"):
    wb = Workbook()
    wb.remove(wb.active)
    _sayfa(
        wb, "PSÇ", "A-B", "PAZARTESİ/SALI/ÇARŞAMBA ",
        ogrenciler=[
            ("KAMiL  DENEMEOĞLU", SARI),      # küçük i + çift boşluk
            ("SELiN TESTOĞLU ", SARI),        # son boşluk
            ("DENEME ÖĞRENCİ k", SARI),       # ad sonunda cinsiyet
            ("ÖRNEK KİŞİ", YESIL),
            ("TEST ADI", YESIL),
        ],
        yerler=["1.KAT CERR.", "4.KAT DAH.", "ALÇI ODASI e", "DİŞ", "DİŞ"],
        kurallar=[
            ("1.KAT CERR.", "1 ÖĞRENCİ", ["CERRAHİ BRANŞ"], False),
            ("4.KAT DAH.", "1 ÖĞRENCİ", ["DAHİLİ BRANŞ"], False),
            ("6.KAT CERR.", "2 ÖĞRENCİ", ["YOĞUN BÖLÜM", "CERRAHİ BRANŞ"], False),
            ("ALÇI ODASI e", "1 ERKEK ÖĞRENCİ", ["HAFİF BÖLÜM"], True),
            ("DİŞ", "BİRİNCİ KİŞİ, İKİNCİ KİŞİ", ["HAFİF BÖLÜM"], False),
            ("ECZANE  SARF e", "BU GRUPTA ÖĞRENCİ VERME", [], True),
        ],
        kural_basligi=kural_basligi,
    )
    _sayfa(
        wb, "ÇPC", "C-D", "ÇARŞAMBA/PERŞEMBE/CUMA ",
        ogrenciler=[("C SINIFI BİR", MAVI), ("D SINIFI BİR", PEMBE),
                    ("D SINIFI İKİ", PEMBE)],
        yerler=["ECZ. SARF e", "DOĞ.SAL-Y.DOĞ", "EKG"],
        kurallar=[
            ("ECZANE  SARF e", "1 ERKEK ÖĞRENCİ", ["HAFİF BÖLÜM"], True),
            ("DOĞ.SAL-Y.DOĞ k", "1 KIZ ÖĞRENCİ", ["HAFİF BÖLÜM"], True),
            ("EKG k", "1 KIZ ÖĞRENCİ", ["HAFİF BÖLÜM"], True),
        ],
        kural_basligi=kural_basligi,
    )
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def sonuc():
    return read_pairing_workbook(_kitap())


# ---- yardımcı ayrıştırıcılar ----

@pytest.mark.parametrize("metin,beklenen", [
    ("1 ÖĞRENCİ", 1), ("2 ÖĞRENCİ", 2), ("1 KIZ ÖĞRENCİ", 1),
    ("1 ERKEK ÖĞRENCİ", 1), ("BU GRUPTA ÖĞRENCİ VERME", 0),
    ("ÜÇÜNCÜ KİŞİ, DÖRDÜNCÜ KİŞİ", 2),
    (" BİRİNCİ KİŞİ, İKİNCİ KİŞİ ocak a kadar", 2),
    ("", None), (None, None),
])
def test_quota_text(metin, beklenen):
    assert quota_from_text(metin) == beklenen


@pytest.mark.parametrize("ham,ad,cinsiyet", [
    ("KAMiL  DENEMEOĞLU", "KAMİL DENEMEOĞLU", ""),
    ("SELiN TESTOĞLU ", "SELİN TESTOĞLU", ""),
    ("NUR  ÖRNEKOĞLU k", "NUR ÖRNEKOĞLU", "K"),
    ("KİŞİ DENEMEOĞLU e", "KİŞİ DENEMEOĞLU", "E"),
])
def test_student_name_cleanup(ham, ad, cinsiyet):
    assert clean_student_name(ham) == (ad, cinsiyet)


@pytest.mark.parametrize("ham,kanonik,ek", [
    ("4.KAT DAH.", "4.KAT DAHİLİYE", ""),
    ("2.KAT DOĞUM SERVİSİ k", "2.KAT DOĞUM", "K"),
    ("DOĞ.SAL-Y.DOĞ k", "DOĞUM SALONU-YENİ DOĞAN", "K"),
    ("ECZ. SARF e", "ECZANE SARF", "E"),
    ("ECZANE  SARF e", "ECZANE SARF", "E"),
    ("FTR ÜNİT", "FTR.ÜNİT", ""),
    ("KÜÇÜK CERR ODASI k", "KÜÇÜK CERRAHİ ODASI", "K"),
])
def test_abbreviations_resolve_to_catalog_names(ham, kanonik, ek):
    """Gereksinimler birimleri ADLA gösteriyor; "4.KAT DAH" diye içe aktarılan
    birim "4.KAT DAHİLİYE" isteyen gereksinime sessizce uymazdı."""
    assert canonical_unit_name(ham) == (kanonik, ek)


def test_canonical_names_exist_in_the_default_catalog():
    from rotasyon.catalog import load_default_units
    from rotasyon.normalize import normalize_key
    from rotasyon.pairing_import import CANONICAL

    katalog = {normalize_key(u.name) for u in load_default_units()}
    for hedef in CANONICAL.values():
        assert normalize_key(hedef) in katalog, hedef


# ---- algılama ----

def test_detects_pairing_workbook():
    assert is_pairing_workbook(_kitap())


def test_plain_rotation_is_not_mistaken_for_pairing():
    wb = Workbook()
    ws = wb.active
    for j, ay in enumerate(AYLAR, start=2):
        ws.cell(1, j, ay)
    ws.cell(2, 1, "BİR ÖĞRENCİ")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    assert not is_pairing_workbook(buf)


# ---- öğrenciler ----

def test_reads_only_students_not_the_rules_block(sonuc):
    """Genel okuyucu kural bloğunu da öğrenci sanıyordu (61 → 130)."""
    assert len(sonuc.students) == 8
    adlar = {s.full_name for s in sonuc.students}
    assert not any("ÖĞRENCİ VERME" in a or "KAT" in a or "AYLIK" in a for a in adlar)


def test_rules_block_without_the_word_kurallari_still_stops_the_list():
    """Bir belgede başlıkta "KURALLARI" yazmıyordu; başlık öğrenci sanılıyordu."""
    r = read_pairing_workbook(_kitap(kural_basligi=""))
    assert len(r.students) == 8


def test_groups_follow_sheet_order(sonuc):
    assert sum(1 for s in sonuc.students if s.group == 1) == 5
    assert sum(1 for s in sonuc.students if s.group == 2) == 3


def test_class_letters_come_from_the_pairing_header(sonuc):
    """Renk→harf sabit değil (A sarısı bir dosyada FFFF99, ötekinde FFF2CC);
    harf, başlıktaki eşleşmeden ve renklerin görülme sırasından gelir."""
    sayim = {(c.name, c.group): c.count for c in sonuc.classes}
    assert sayim == {("A", 1): 3, ("B", 1): 2, ("C", 2): 1, ("D", 2): 2}


def test_class_colours_are_kept(sonuc):
    renk = {c.name: c.color for c in sonuc.classes}
    assert renk["A"] == f"#{SARI}" and renk["D"] == f"#{PEMBE}"


def test_gender_suffix_on_names_is_read(sonuc):
    k = [s for s in sonuc.students if s.gender == "K"]
    assert len(k) == 1 and k[0].full_name == "DENEME ÖĞRENCİ"


def test_names_are_cleaned(sonuc):
    adlar = {s.full_name for s in sonuc.students}
    assert "KAMİL DENEMEOĞLU" in adlar
    assert "SELİN TESTOĞLU" in adlar


def test_day_labels(sonuc):
    assert "PAZARTESİ" in sonuc.day_labels[1]
    assert "CUMA" in sonuc.day_labels[2]


def test_mismatched_colours_warn_instead_of_guessing():
    wb = Workbook()
    wb.remove(wb.active)
    _sayfa(wb, "PSÇ", "A-B", "PZT",
           ogrenciler=[("X Y", SARI), ("Z W", YESIL), ("Q R", PEMBE)],
           yerler=["ACİL"] * 3, kurallar=[("ACİL", "3 ÖĞRENCİ", [], False)])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = read_pairing_workbook(buf)
    assert all(not s.class_name for s in r.students)
    assert any("ad rengi" in w for w in r.warnings)


# ---- birimler ----

def _birim(sonuc, grup, ad):
    return next(u for u in sonuc.units if u.group == grup and u.name == ad)


def test_quota_gender_branch_intensity(sonuc):
    u = _birim(sonuc, 1, "6.KAT CERR")
    assert (u.quota, u.branch, u.intensity) == (2, BRANCH_CERRAHI, INTENSITY_YOGUN)
    u = _birim(sonuc, 1, "4.KAT DAHİLİYE")
    assert (u.quota, u.branch) == (1, BRANCH_DAHILI)
    u = _birim(sonuc, 1, "ALÇI ODASI")
    assert (u.quota, u.gender, u.intensity) == (1, "E", INTENSITY_HAFIF)


def test_quota_written_as_names_uses_the_slot_list(sonuc):
    """DİŞ kontenjanı yerine iki kişinin adı yazılmış → 2."""
    assert _birim(sonuc, 1, "DİŞ").quota == 2


def test_closed_in_this_group(sonuc):
    assert _birim(sonuc, 1, "ECZANE SARF").quota == 0
    assert _birim(sonuc, 2, "ECZANE SARF").quota == 1


def test_gender_from_rules_block_beats_the_slot_list(sonuc):
    """Yer çizelgesinde "DOĞ.SAL-Y.DOĞ" eksiz yazılmış; kural bloğunda "k" var."""
    assert _birim(sonuc, 2, "DOĞUM SALONU-YENİ DOĞAN").gender == "K"
    assert _birim(sonuc, 2, "EKG").gender == "K"


def test_rule_text_lines_are_not_units(sonuc):
    adlar = {u.name for u in sonuc.units}
    assert not any("ÖĞRENCİ" in a or "BRANŞ" in a or "AYLIK" in a for a in adlar)


def test_unit_only_in_rules_block_is_reported(sonuc):
    """6.KAT CERR kural bloğunda 2 kişi ama yer çizelgesinde yok."""
    assert any("6.KAT CERR" in w and "yer çizelgesinde hiç yok" in w
               for w in sonuc.warnings)


# ---- katalogla birleştirme ----

def test_catalog_properties_not_in_the_document_are_kept(sonuc):
    """Belge doldurma politikası/ebelik önceliği yazmıyor; katalogdaki kalır."""
    katalog = [
        Unit(name="DOĞUM SALONU-YENİ DOĞAN", group=2, quota=5,
             ebelik_priority=True, single_slot=True),
    ]
    birimler, _ = build_units(sonuc, katalog)
    u = next(x for x in birimler if x.group == 2 and x.name == "DOĞUM SALONU-YENİ DOĞAN")
    assert u.ebelik_priority and u.single_slot
    assert u.quota == 1 and u.gender_restriction == "K"  # belgeden


def test_catalog_units_missing_from_the_document_become_inactive(sonuc):
    katalog = [Unit(name="ÜROLOJİ POLK", group=1, quota=1, active=True)]
    birimler, notlar = build_units(sonuc, katalog)
    u = next(x for x in birimler if x.name == "ÜROLOJİ POLK")
    assert not u.active
    assert any("ÜROLOJİ POLK" in n for n in notlar)


def test_closed_unit_keeps_a_usable_quota(sonuc):
    """Pasif yapılan birimin kontenjanı 0 değil 1: açılınca kullanılabilsin."""
    birimler, _ = build_units(sonuc, [])
    u = next(x for x in birimler if x.group == 1 and x.name == "ECZANE SARF")
    assert not u.active and u.quota == 1


def test_imported_units_satisfy_default_requirements(sonuc):
    """Belgeden gelen adlar varsayılan gereksinimlerin istediği adlarla
    eşleşmeli — yoksa gereksinim sessizce hiçbir birime uymaz."""
    from rotasyon.catalog import load_default_units
    from rotasyon.normalize import normalize_key

    birimler, _ = build_units(sonuc, load_default_units())
    adlar = {normalize_key(u.name) for u in birimler}
    assert normalize_key("4.KAT DAHİLİYE") in adlar
    assert normalize_key("DOĞUM SALONU-YENİ DOĞAN") in adlar
    assert normalize_key("4.KAT DAH") not in adlar


# ---- öğrenci modelleri ----

def test_programs_are_applied_per_class(sonuc):
    ogr, siniflar = build_students(sonuc, {"D": PROGRAM_EBELIK})
    assert {s.program for s in ogr if s.class_name == "D"} == {PROGRAM_EBELIK}
    assert {s.program for s in ogr if s.class_name == "A"} == {PROGRAM_HEMSIRELIK}
    assert [c.name for c in siniflar] == ["A", "B", "C", "D"]


# ---- gerçek belgeler (varsa; yalnız sayılar) ----

ORNEKLER = {
    "AB-CD": ({1: 36, 2: 28}, {"A": 18, "B": 18, "C": 20, "D": 8}),
    "AC-BD": ({1: 38, 2: 26}, {"A": 18, "B": 18, "C": 20, "D": 8}),
    "AD-BC": ({1: 26, 2: 38}, {"A": 18, "B": 18, "C": 20, "D": 8}),
}


@pytest.mark.parametrize("ad", sorted(ORNEKLER))
def test_real_documents(ad):
    yol = Path("example") / f"{ad} ROTASYON.xlsx"
    if not yol.exists():
        pytest.skip("gerçek belge yok (git dışı)")
    gruplar, siniflar = ORNEKLER[ad]
    r = read_pairing_workbook(yol)
    assert {g: sum(1 for s in r.students if s.group == g) for g in (1, 2)} == gruplar
    assert {c.name: c.count for c in r.classes} == siniflar
    assert not [s for s in r.students if not s.class_name]
    assert {g: sum(1 for u in r.units if u.group == g) for g in (1, 2)} == {1: 28, 2: 28}


# ---- BÖLÜMLER listesinde gri = o grupta iptal ----

def _gri_kitap(gri_renk="A6A6A6", blok_kota="1 ERKEK ÖĞRENCİ", yerde=False):
    """ALÇI ODASI sağdaki BÖLÜMLER listesinde griye boyanmış bir sayfa."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = _sayfa(
        wb, "ÇPC", "C-D", "ÇARŞAMBA/PERŞEMBE/CUMA",
        ogrenciler=[("BİR KİŞİ", MAVI), ("İKİ KİŞİ", PEMBE)],
        yerler=["ACİL", "ALÇI ODASI e"] if yerde else ["ACİL", "ACİL"],
        kurallar=[
            ("ACİL", "2 ÖĞRENCİ", ["YOĞUN BÖLÜM"], False),
            ("ALÇI ODASI e", blok_kota, ["HAFİF BÖLÜM"], True),
        ],
    )
    for i, ad in enumerate(["ACİL", "ALÇI ODASI"], start=4):
        c = ws.cell(i, 16, ad)
        if ad == "ALÇI ODASI" and gri_renk:
            c.fill = _dolgu(gri_renk)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return read_pairing_workbook(buf)


@pytest.mark.parametrize("renk,beklenen", [
    ("#A6A6A6", True), ("#BFBFBF", True), ("#808080", True), ("#D9D9D9", True),
    ("#F2F2F2", False),  # zemin süsü, iptal değil
    ("#FFFFFF", False), ("#000000", False),
    ("#92D050", False), ("#FFF2CC", False), ("", False),
])
def test_grey_detection(renk, beklenen):
    from rotasyon.pairing_import import is_grey

    assert is_grey(renk) is beklenen


def test_grey_in_master_list_closes_the_unit_even_if_the_block_says_one():
    """Gerçek belgede oldu: birim griye boyanıp yer çizelgesinden çıkarılmış,
    bloktaki "1 ERKEK ÖĞRENCİ" unutulmuş. Gri, açık bir iptal kararıdır."""
    r = _gri_kitap()
    alci = next(u for u in r.units if u.name == "ALÇI ODASI")
    assert alci.quota == 0
    assert alci.gender == "E" and alci.intensity == INTENSITY_HAFIF  # kısıtları korunur
    assert any("griye boyanmış" in w and "KAPALI" in w for w in r.warnings)


def test_grey_and_verme_agree_without_a_warning():
    r = _gri_kitap(blok_kota="BU GRUPTA ÖĞRENCİ VERME")
    assert next(u for u in r.units if u.name == "ALÇI ODASI").quota == 0
    assert not any("griye boyanmış" in w for w in r.warnings)


def test_no_grey_keeps_the_block_quota():
    """Karşıt kanıt: boyanmamış birim bloğun dediği gibi açık kalır."""
    r = _gri_kitap(gri_renk="")
    assert next(u for u in r.units if u.name == "ALÇI ODASI").quota == 1


def test_light_grey_background_is_not_a_cancellation():
    r = _gri_kitap(gri_renk="F2F2F2")
    assert next(u for u in r.units if u.name == "ALÇI ODASI").quota == 1


def test_grey_wins_even_if_the_slot_list_still_has_it():
    r = _gri_kitap(yerde=True)
    assert next(u for u in r.units if u.name == "ALÇI ODASI").quota == 0


def test_closed_by_grey_is_inactive_after_catalog_merge():
    birimler, _ = build_units(_gri_kitap(), [])
    alci = next(u for u in birimler if u.name == "ALÇI ODASI")
    assert not alci.active and alci.quota == 1 and alci.gender_restriction == "E"
